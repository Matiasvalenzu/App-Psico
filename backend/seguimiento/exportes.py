"""Exportación PDF/Word de brief e informes, con encabezado profesional y bloque de firma."""
from io import BytesIO

from django.utils import timezone


def _encabezado(doc):
    u = doc.psicologo
    perfil = getattr(u, "perfil_psicologo", None)
    lineas = [u.get_full_name() or u.username]
    if perfil:
        if perfil.especialidad_clinica:
            lineas.append(perfil.especialidad_clinica)
        if perfil.rut_profesional:
            lineas.append(f"RUT: {perfil.rut_profesional}")
        if perfil.registro_profesional:
            lineas.append(f"Registro profesional: {perfil.registro_profesional}")
        if perfil.comuna:
            lineas.append(perfil.comuna)
    return lineas


def _titulo(doc):
    return doc.get_tipo_display() + (" (BORRADOR)" if doc.es_borrador else "")


def _firma(doc):
    return ["", "", "_______________________________", doc.psicologo.get_full_name() or "", "Firma del profesional"]


def pdf_documento(doc):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    buf = BytesIO()
    styles = getSampleStyleSheet()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                            title=doc.get_tipo_display())
    story = [Paragraph(escape(l), styles["Normal"]) for l in _encabezado(doc)]
    story += [Spacer(1, 12), Paragraph(escape(_titulo(doc)), styles["Title"]),
              Paragraph(escape(f"Fecha: {timezone.localtime(doc.updated_at).strftime('%d-%m-%Y')}"), styles["Normal"])]
    if doc.destinatario:
        story.append(Paragraph(escape(f"Destinatario: {doc.destinatario}"), styles["Normal"]))
    story.append(Spacer(1, 12))
    for linea in (doc.contenido or "").splitlines():
        if not linea.strip():
            story.append(Spacer(1, 6))
        elif linea.strip().isupper() and len(linea) < 90:
            story.append(Paragraph(f"<b>{escape(linea.strip())}</b>", styles["Heading4"]))
        else:
            story.append(Paragraph(escape(linea), styles["BodyText"]))
    story += [Paragraph(escape(l) or "&nbsp;", styles["Normal"]) for l in _firma(doc)]
    pdf.build(story)
    return buf.getvalue()


def docx_documento(doc):
    from docx import Document

    d = Document()
    for l in _encabezado(doc):
        d.add_paragraph(l)
    d.add_heading(_titulo(doc), level=1)
    d.add_paragraph(f"Fecha: {timezone.localtime(doc.updated_at).strftime('%d-%m-%Y')}")
    if doc.destinatario:
        d.add_paragraph(f"Destinatario: {doc.destinatario}")
    for linea in (doc.contenido or "").splitlines():
        if linea.strip().isupper() and 0 < len(linea) < 90:
            d.add_heading(linea.strip(), level=2)
        else:
            d.add_paragraph(linea)
    for l in _firma(doc):
        d.add_paragraph(l)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()
