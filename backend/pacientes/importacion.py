import csv
import io
import json
import re
import unicodedata
from datetime import date, datetime

from django.utils import timezone
from rest_framework import serializers

from .documentos import TIPO_RUT, formatear_documento, normalizar_documento

MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 1000
CAMPOS = (
    "nombre", "apellido", "rut", "fecha_nacimiento", "edad",
    "email", "telefono", "direccion", "comuna", "prevision",
)
SINONIMOS = {
    "nombre": ("nombre", "nombres", "name", "primer nombre"),
    "apellido": ("apellido", "apellidos", "apellido paterno"),
    "rut": ("rut", "run", "documento", "dni"),
    "fecha_nacimiento": ("fecha nacimiento", "fecha de nacimiento", "nacimiento", "fnac"),
    "edad": ("edad", "anos", "años"),
    "email": ("email", "correo", "e-mail", "mail", "correo electronico"),
    "telefono": ("telefono", "teléfono", "celular", "whatsapp", "movil", "móvil", "fono"),
    "direccion": ("direccion", "dirección", "domicilio", "calle"),
    "comuna": ("comuna", "ciudad"),
    "prevision": ("prevision", "previsión", "isapre", "fonasa"),
}
CONTACTO = ("telefono_whatsapp", "email_contacto", "direccion", "comuna", "prevision", "fecha_nacimiento")


class ArchivoImportacionInvalido(Exception):
    pass


def _clave(texto):
    limpio = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", limpio.lower()).strip()


def _parece_email(valor):
    return bool(re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", str(valor or "").strip()))


def _parece_rut(valor):
    try:
        normalizar_documento(TIPO_RUT, str(valor or ""))
    except serializers.ValidationError:
        return False
    return True


def _parece_telefono(valor):
    digitos = re.sub(r"\D", "", str(valor or ""))
    return (len(digitos) == 9 and digitos.startswith("9")) or (
        digitos.startswith("56") and len(digitos) == 11 and digitos[2] == "9"
    )


def sugerir_mapeo(encabezados, filas):
    usados = set()
    mapeo = {}
    for indice, encabezado in enumerate(encabezados):
        clave = _clave(encabezado)
        campo = next((c for c, nombres in SINONIMOS.items() if clave in {_clave(n) for n in nombres} and c not in usados), "")
        if not campo:
            muestra = [fila[indice] for fila in filas if indice < len(fila) and str(fila[indice] or "").strip()]
            if muestra and campo_libre("rut", usados) and sum(_parece_rut(v) for v in muestra) >= max(1, len(muestra) // 2):
                campo = "rut"
            elif muestra and campo_libre("email", usados) and sum(_parece_email(v) for v in muestra) >= max(1, len(muestra) // 2):
                campo = "email"
            elif muestra and campo_libre("telefono", usados) and sum(_parece_telefono(v) for v in muestra) >= max(1, len(muestra) // 2):
                campo = "telefono"
        if campo:
            usados.add(campo)
            mapeo[str(indice)] = campo
    return mapeo


def campo_libre(campo, usados):
    return campo not in usados


def leer_archivo(nombre, datos):
    if len(datos) > MAX_BYTES:
        raise ArchivoImportacionInvalido("El archivo supera los 5 MB.")
    extension = (nombre or "").rsplit(".", 1)[-1].lower()
    try:
        if extension == "csv":
            return _leer_csv(datos)
        if extension == "xlsx":
            return _leer_xlsx(datos)
        if extension == "xls":
            return _leer_xls(datos)
    except ArchivoImportacionInvalido:
        raise
    except Exception as exc:
        raise ArchivoImportacionInvalido("No se pudo leer el archivo. Usa la plantilla .xlsx o un CSV.") from exc
    raise ArchivoImportacionInvalido("Formato no soportado. Usa .xlsx, .xls o .csv.")


def _filas_limitadas(filas):
    utiles = [fila for fila in filas if any(str(c or "").strip() for c in fila)]
    if not utiles:
        raise ArchivoImportacionInvalido("El archivo no tiene filas.")
    if len(utiles) - 1 > MAX_FILAS:
        raise ArchivoImportacionInvalido(f"El archivo tiene más de {MAX_FILAS} filas.")
    encabezados = [str(c or "").strip() for c in utiles[0]]
    if not any(encabezados):
        raise ArchivoImportacionInvalido("La primera fila debe tener los nombres de columna.")
    return encabezados, utiles[1:]


def _leer_csv(datos):
    texto = None
    for codec in ("utf-8-sig", "cp1252"):
        try:
            texto = datos.decode(codec)
            break
        except UnicodeDecodeError:
            continue
    if texto is None:
        raise ArchivoImportacionInvalido("No se pudo leer el CSV. Guárdalo en UTF-8.")
    muestra = texto[:4096]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t")
    except csv.Error:
        dialecto = csv.excel
    lector = csv.reader(io.StringIO(texto), dialecto)
    return _filas_limitadas(list(lector))


def _leer_xlsx(datos):
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    try:
        hoja = libro.active
        return _filas_limitadas([list(fila) for fila in hoja.iter_rows(values_only=True)])
    finally:
        libro.close()


def _leer_xls(datos):
    import xlrd

    libro = xlrd.open_workbook(file_contents=datos)
    hoja = libro.sheet_by_index(0)
    filas = []
    for i in range(hoja.nrows):
        fila = []
        for j in range(hoja.ncols):
            valor = hoja.cell_value(i, j)
            if hoja.cell_type(i, j) == xlrd.XL_CELL_DATE:
                valor = xlrd.xldate_as_datetime(valor, libro.datemode).date()
            fila.append(valor)
        filas.append(fila)
    return _filas_limitadas(filas)


def _fecha(valor):
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()[:10]
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    return None


def _telefono(valor):
    digitos = re.sub(r"\D", "", str(valor or ""))
    if not digitos:
        return "", None
    if digitos.startswith("56") and len(digitos) == 11 and digitos[2] == "9":
        return f"+{digitos}", None
    if len(digitos) == 9 and digitos.startswith("9"):
        return f"+56{digitos}", None
    return "", "Teléfono no reconocido como móvil chileno; no se importó."


def _email(valor):
    texto = str(valor or "").strip().lower()
    if not texto:
        return "", None
    if not _parece_email(texto):
        return "", "Correo inválido; no se importó."
    return texto, None


def _celda(fila, indice):
    if indice is None or indice >= len(fila):
        return ""
    valor = fila[indice]
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return valor


def columnas_de(mapeo):
    por_campo = {}
    for indice, campo in (mapeo or {}).items():
        if campo in CAMPOS and campo not in por_campo:
            por_campo[campo] = int(indice)
    return por_campo


def evaluar_filas(encabezados, filas, mapeo, existentes):
    columnas = columnas_de(mapeo)
    mapeados = {int(k) for k, v in (mapeo or {}).items() if str(v) in CAMPOS}
    ignoradas = [h for i, h in enumerate(encabezados) if i not in mapeados and h]
    resultado = []
    vistos_rut = {}
    vistos_email = {}
    for numero, fila in enumerate(filas, start=1):
        motivos = []
        estado = "ok"
        datos = _datos_fila(fila, columnas, motivos)
        if datos is None:
            resultado.append({"fila": numero, "estado": "error", "motivos": motivos, "datos": {}, "duplicado": None})
            continue
        duplicado = _duplicado(datos, existentes, vistos_rut, vistos_email, numero)
        if duplicado:
            estado = "error" if duplicado["tipo"] == "conflicto" else "advertencia"
            motivos.append(duplicado["motivo"])
        if motivos and estado == "ok":
            estado = "advertencia"
        if datos.get("rut"):
            vistos_rut[datos["rut"]] = numero
        if datos.get("email"):
            vistos_email[datos["email"]] = numero
        resultado.append({"fila": numero, "estado": estado, "motivos": motivos, "datos": datos, "duplicado": duplicado})
    return {"columnas_ignoradas": [c for c in ignoradas if c], "filas": resultado}


def _datos_fila(fila, columnas, motivos):
    def valor(campo):
        return _celda(fila, columnas.get(campo))

    nombre = str(valor("nombre") or "").strip()
    apellido = str(valor("apellido") or "").strip()
    if not nombre and not apellido:
        motivos.append("Falta el nombre.")
        return None
    if not apellido and " " in nombre:
        nombre, apellido = nombre.rsplit(" ", 1)
        motivos.append("El apellido se tomó de la última palabra del nombre.")
    if not nombre or not apellido:
        motivos.append("Faltan nombre y apellido.")
        return None
    if len(nombre) > 100 or len(apellido) > 100:
        motivos.append("El nombre o el apellido superan los 100 caracteres.")
        return None

    rut = ""
    rut_texto = str(valor("rut") or "").strip()
    if rut_texto:
        try:
            rut = normalizar_documento(TIPO_RUT, rut_texto)
        except serializers.ValidationError:
            motivos.append("RUT inválido.")
            return None

    fecha = _fecha(valor("fecha_nacimiento"))
    if columnas.get("fecha_nacimiento") is not None and str(valor("fecha_nacimiento") or "").strip() and not fecha:
        motivos.append("La fecha de nacimiento no se entendió.")
    edad_txt = str(valor("edad") or "").strip()
    edad = int(edad_txt) if edad_txt.isdigit() else None
    if not fecha and edad and 0 <= edad <= 120:
        fecha = date(timezone.localdate().year - edad, 1, 1)
        motivos.append("La fecha de nacimiento se estimó desde la edad (1 de enero).")
    email, aviso = _email(valor("email"))
    if aviso:
        motivos.append(aviso)
    telefono, aviso = _telefono(valor("telefono"))
    if aviso:
        motivos.append(aviso)
    return {
        "nombre": nombre,
        "apellido": apellido,
        "rut": rut,
        "fecha_nacimiento": fecha.isoformat() if fecha else "",
        "edad": edad,
        "email": email,
        "telefono": telefono,
        "direccion": str(valor("direccion") or "").strip()[:255],
        "comuna": str(valor("comuna") or "").strip()[:100],
        "prevision": str(valor("prevision") or "").strip()[:100],
    }


def _duplicado(datos, existentes, vistos_rut, vistos_email, numero):
    if datos["rut"] and datos["rut"] in vistos_rut:
        return {"tipo": "archivo", "fila": vistos_rut[datos["rut"]], "motivo": "RUT repetido en el archivo."}
    if datos["email"] and datos["email"] in vistos_email:
        return {"tipo": "archivo", "fila": vistos_email[datos["email"]], "motivo": "Correo repetido en el archivo."}
    por_rut = existentes["rut"].get(datos["rut"]) if datos["rut"] else None
    por_email = existentes["email"].get(datos["email"]) if datos["email"] else None
    if por_rut and por_email and por_rut != por_email:
        return {"tipo": "conflicto", "paciente": None, "motivo": "El RUT y el correo coinciden con pacientes distintos."}
    paciente = por_rut or por_email
    if not paciente:
        return None
    via = "RUT" if por_rut else "correo"
    return {"tipo": "existente", "paciente": paciente, "motivo": f"Ya existe un paciente con este {via}."}


def indice_existentes(pacientes):
    rut, email = {}, {}
    for paciente in pacientes:
        if paciente.documento_normalizado:
            rut[paciente.documento_normalizado] = paciente.id
        if paciente.email_contacto:
            email[paciente.email_contacto.strip().lower()] = paciente.id
    return {"rut": rut, "email": email}


def aplicar(filas, decisiones, usuario):
    from .models import Paciente

    creados, actualizados, omitidos = [], [], []
    anteriores = []
    errores = []
    for fila in filas:
        if fila["estado"] == "error":
            omitidos.append(fila["fila"])
            errores.append({"fila": fila["fila"], "motivo": "; ".join(fila["motivos"])})
            continue
        duplicado = fila["duplicado"]
        if duplicado and duplicado["tipo"] == "conflicto":
            omitidos.append(fila["fila"])
            errores.append({"fila": fila["fila"], "motivo": duplicado["motivo"]})
            continue
        if duplicado:
            decision = (decisiones or {}).get(str(fila["fila"]), "")
            if decision not in ("saltar", "actualizar"):
                raise ArchivoImportacionInvalido(f"Elige saltar o actualizar la fila {fila['fila']}.")
            if decision == "saltar":
                omitidos.append(fila["fila"])
                continue
            paciente_id = duplicado.get("paciente")
            if duplicado["tipo"] == "archivo":
                paciente_id = _id_de_fila(creados, actualizados, duplicado["fila"])
            if not paciente_id:
                omitidos.append(fila["fila"])
                errores.append({"fila": fila["fila"], "motivo": "No se encontró el paciente a actualizar."})
                continue
            paciente = Paciente.objects.get(pk=paciente_id, psicologo=usuario)
            anteriores.append(_snapshot(paciente))
            _actualizar_contacto(paciente, fila["datos"])
            actualizados.append({"fila": fila["fila"], "paciente": paciente.id})
            continue
        paciente = _crear(fila["datos"], usuario)
        creados.append({"fila": fila["fila"], "paciente": paciente.id})
    return creados, actualizados, omitidos, anteriores, errores


def _id_de_fila(creados, actualizados, fila):
    for grupo in (creados, actualizados):
        for item in grupo:
            if item["fila"] == fila:
                return item["paciente"]
    return None


def _snapshot(paciente):
    datos = {"id": paciente.id}
    for campo in CONTACTO:
        valor = getattr(paciente, campo)
        datos[campo] = valor.isoformat() if isinstance(valor, date) else valor
    return datos


def _actualizar_contacto(paciente, datos):
    paciente.telefono_whatsapp = datos["telefono"] or paciente.telefono_whatsapp
    paciente.email_contacto = datos["email"] or paciente.email_contacto
    paciente.direccion = datos["direccion"] or paciente.direccion
    paciente.comuna = datos["comuna"] or paciente.comuna
    paciente.prevision = datos["prevision"] or paciente.prevision
    if datos["fecha_nacimiento"] and not paciente.fecha_nacimiento:
        paciente.fecha_nacimiento = datos["fecha_nacimiento"]
    paciente.save(update_fields=[*CONTACTO, "updated_at"])


def _crear(datos, usuario):
    from .models import Paciente

    rut = formatear_documento(TIPO_RUT, datos["rut"]) if datos["rut"] else ""
    return Paciente.objects.create(
        psicologo=usuario,
        nombre=datos["nombre"],
        apellido=datos["apellido"],
        tipo_documento=TIPO_RUT,
        numero_documento=rut,
        documento_normalizado=datos["rut"],
        rut=rut,
        fecha_nacimiento=datos["fecha_nacimiento"] or None,
        edad=datos["edad"],
        telefono_whatsapp=datos["telefono"],
        email_contacto=datos["email"],
        direccion=datos["direccion"],
        comuna=datos["comuna"],
        prevision=datos["prevision"],
    )


def errores_csv(errores):
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(["fila", "motivo"])
    for error in errores:
        writer.writerow([error["fila"], error["motivo"]])
    return buffer.getvalue()


def plantilla_xlsx():
    from openpyxl import Workbook

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Pacientes"
    hoja.append(["nombre", "apellido", "rut", "fecha nacimiento", "edad", "email", "teléfono", "dirección", "comuna", "previsión"])
    hoja.append(["Ana", "Rojas", "11.111.111-1", "13/10/1990", "", "ana.rojas@example.com", "987654321", "Calle Falsa 123", "Ñuñoa", "Fonasa"])
    hoja.append(["Luis", "Pérez", "", "", "34", "luis.perez@example.com", "+56 9 1111 2222", "", "Providencia", ""])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def sugerir_con_ia(encabezados, muestra):
    import requests
    from django.conf import settings

    if not settings.DEEPSEEK_API_KEY:
        raise ArchivoImportacionInvalido("La sugerencia con IA no está configurada.")
    recorte = [[str(c or "")[:80] for c in fila[:20]] for fila in muestra[:3]]
    encabezados = [str(c or "")[:80] for c in encabezados[:20]]
    prompt = (
        "Asigna columnas de una planilla de pacientes a estos campos y a ningún otro: "
        + ", ".join(CAMPOS)
        + ". Responde solo JSON {\"indice\": \"campo\"}. Encabezados: "
        + json.dumps(encabezados, ensure_ascii=False)
        + ". Muestra de 3 filas: "
        + json.dumps(recorte, ensure_ascii=False)
    )
    response = requests.post(
        f"{settings.DEEPSEEK_BASE_URL}/chat/completions",
        headers={"Authorization": f"Bearer {settings.DEEPSEEK_API_KEY}", "Content-Type": "application/json"},
        json={
            "model": "deepseek-chat",
            "messages": [{"role": "user", "content": prompt}],
            "temperature": 0,
            "max_tokens": 400,
        },
        timeout=30,
    )
    response.raise_for_status()
    texto = response.json()["choices"][0]["message"]["content"]
    inicio, fin = texto.find("{"), texto.rfind("}")
    if inicio < 0 or fin < inicio:
        raise ArchivoImportacionInvalido("La IA no devolvió un mapeo usable.")
    crudo = json.loads(texto[inicio : fin + 1])
    return {str(k): v for k, v in crudo.items() if v in CAMPOS}
